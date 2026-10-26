"""Rebuild docs/sources-tracker.xlsx from docs/sources-tracker.csv.

    pip install openpyxl
    python pipeline/tools/tracker_xlsx.py

The CSV is the source of truth: the admin page's Sources tab reads it and saves edits back to
it. The spreadsheet is the same table for people who prefer Excel, with the dropdowns and a
summary sheet. Regenerate it after editing the CSV; do not edit both.
"""
from __future__ import annotations
import csv, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CSV = ROOT / "docs" / "sources-tracker.csv"
XLSX = ROOT / "docs" / "sources-tracker.xlsx"

STATUSES = ["In use", "Check", "Disabled", "Candidate", "Declined"]
ROUTES = ["Tribe REST", "CivicPlus iCal", "Plain iCal", "Squarespace JSON", "HTML listing", "WhoFi", "LibCal",
          "Event plugin API", "Eventbrite", "Wix (ask)", "Ticketing (ask)", "Submission form", "Check source", "None"]
PRIORITIES = ["High", "Medium", "Low"]
WAVES = ["1", "2", "3", "4"]
ACTIONS = ["Add now", "Enable", "Build connector", "Write selectors", "Email for feed address", "Ask for export",
           "Offer the form", "Confirm domain", "Retry probe", "Keep as check", "Declined", "Done"]
RESPONSES = ["Not contacted", "No reply", "Replied", "Feed received", "Declined"]
YESNO = ["Yes", "No"]
LISTS = {"Status": STATUSES, "Route": ROUTES, "Priority": PRIORITIES, "Wave": WAVES, "Action": ACTIONS,
         "Response": RESPONSES, "Live on site": YESNO}
EDITABLE = set(LISTS) | {"Owner", "Contact name", "Contact email", "Contacted on", "Notes"}
WIDTHS = {"Town": 14, "Organisation": 36, "Domain": 32, "Kind": 12, "Status": 11, "Platform found": 34, "Feeds found": 34,
          "Route": 18, "Priority": 10, "Wave": 7, "Action": 24, "Owner": 12, "Contact name": 20, "Contact email": 26,
          "Contacted on": 13, "Response": 15, "Live on site": 11, "Notes": 60}


def main() -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("pip install openpyxl first")
    rows = list(csv.reader(open(CSV, encoding="utf-8")))
    cols, body = rows[0], rows[1:]
    font = "Arial"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sources"
    ws.append(cols)
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(name=font, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="1A1A1A")
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 14)
    for r in body:
        ws.append(r)
    n = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            cell.font = Font(name=font, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=cols[cell.column - 1] in ("Platform found", "Feeds found", "Notes"))
            if cols[cell.column - 1] in EDITABLE:
                cell.fill = PatternFill("solid", fgColor="FFF9C4")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"
    ls = wb.create_sheet("Lists")
    for j, (title, vals) in enumerate(LISTS.items(), start=1):
        ls.cell(row=1, column=j, value=title).font = Font(name=font, bold=True, size=10)
        for i, v in enumerate(vals, start=2):
            ls.cell(row=i, column=j, value=v)
        col = get_column_letter(j)
        dv = DataValidation(type="list", formula1=f"=Lists!${col}$2:${col}${len(vals) + 1}", allow_blank=True)
        ws.add_data_validation(dv)
        if title in cols:
            t = get_column_letter(cols.index(title) + 1)
            dv.add(f"{t}2:{t}{n}")
    sm = wb.create_sheet("Summary", 0)
    sm.column_dimensions["A"].width = 30
    sm["A1"] = "Slack Tide source tracker"
    sm["A1"].font = Font(name=font, bold=True, size=14)
    sm["A2"] = "Generated from docs/sources-tracker.csv, which the admin page's Sources tab edits. Yellow cells are the editable ones."
    sm["A4"] = "By action"
    a = get_column_letter(cols.index("Action") + 1)
    for i, act in enumerate(ACTIONS, start=5):
        sm.cell(row=i, column=1, value=act)
        sm.cell(row=i, column=2, value=f'=COUNTIF(Sources!${a}$2:${a}${n},A{i})')
    wb.save(XLSX)
    print(f"wrote {XLSX}: {len(body)} rows")


if __name__ == "__main__":
    main()
